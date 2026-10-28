from pathlib import Path
import re
from typing import Iterable
from .models import DocumentChunk

SUPPORTED = {".pdf", ".docx", ".xlsx", ".pptx", ".txt", ".md"}

def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n\n".join(f"[Page {i + 1}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages))
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            rows.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None]
                if values:
                    rows.append(" | ".join(values))
        return "\n".join(rows)
    if suffix == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        slides = []
        for i, slide in enumerate(prs.slides, start=1):
            body = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
            if body:
                slides.append(f"[Slide {i}]\n" + "\n".join(body))
        return "\n\n".join(slides)
    raise ValueError(f"Unsupported file type: {suffix}")

def ingest_file(file_or_path) -> list[dict]:
    name = getattr(file_or_path, "name", None) or str(file_or_path)
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type {suffix}. Supported: {', '.join(sorted(SUPPORTED))}")
    if hasattr(file_or_path, "getvalue"):
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(file_or_path.getvalue())
            temp_path = Path(tmp.name)
        try:
            text = _extract_text(temp_path)
        finally:
            temp_path.unlink(missing_ok=True)
    else:
        text = _extract_text(Path(file_or_path))
    return [{"text": text, "source": Path(name).name, "document_id": re.sub(r"[^a-zA-Z0-9_-]", "_", Path(name).stem)}]

def chunk_document(documents: Iterable[dict], chunk_size: int = 1100, overlap: int = 180) -> list[DocumentChunk]:
    if chunk_size < 100 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Require chunk_size >= 100 and 0 <= overlap < chunk_size")
    chunks = []
    for doc in documents:
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", doc.get("text", "")) if p.strip()]
        current = ""
        for para in paragraphs:
            if len(para) > chunk_size:
                if current:
                    chunks.append(DocumentChunk(current, doc["source"], doc["document_id"]))
                    current = ""
                for start in range(0, len(para), chunk_size - overlap):
                    part = para[start:start + chunk_size]
                    if part.strip():
                        chunks.append(DocumentChunk(part, doc["source"], doc["document_id"]))
            elif len(current) + len(para) + 2 <= chunk_size:
                current = f"{current}\n\n{para}".strip()
            else:
                if current:
                    chunks.append(DocumentChunk(current, doc["source"], doc["document_id"]))
                current = para
        if current:
            chunks.append(DocumentChunk(current, doc["source"], doc["document_id"]))
    return chunks
